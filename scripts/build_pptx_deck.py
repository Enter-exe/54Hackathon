"""Build the Kamion Challenge pitch deck as a native .pptx file, mirroring
the content of the HTML slide deck artifact but laid out with PowerPoint
primitives (tables, native charts, shapes) so it's editable in PowerPoint.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

# ---------- palette (mirrors the HTML deck's light theme) ----------
INK = RGBColor(0x17, 0x1A, 0x1D)
INK_DIM = RGBColor(0x5C, 0x66, 0x70)
ACCENT = RGBColor(0xD9, 0x50, 0x0C)
GOOD = RGBColor(0x2F, 0x8F, 0x57)
BAD = RGBColor(0xB2, 0x3A, 0x2C)
LINE = RGBColor(0xD8, 0xDA, 0xD6)
SURFACE_2 = RGBColor(0xF4, 0xF5, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_HEAD = "Bahnschrift"
FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def add_text(slide, text, left, top, width, height, size=14, color=INK, bold=False,
             font=FONT_BODY, align=PP_ALIGN.LEFT, spacing=1.15, upper=False, tracking=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line.upper() if upper else line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = font
        run.font.color.rgb = color
    return box


def add_eyebrow(slide, text, top=Inches(0.55)):
    return add_text(slide, text, Inches(0.7), top, Inches(10), Inches(0.35),
                     size=13, color=ACCENT, bold=True, font=FONT_MONO, upper=True)


def add_headline(slide, text, top=Inches(0.9), size=34, width=Inches(11)):
    return add_text(slide, text, Inches(0.68), top, width, Inches(1.5),
                     size=size, color=INK, bold=True, font=FONT_HEAD, spacing=0.95)


def add_stripe(slide, top):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), top, Inches(0.9), Pt(4))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    return bar


def add_stat_tile(slide, left, top, width, height, value, label):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid(); box.fill.fore_color.rgb = SURFACE_2
    box.line.color.rgb = LINE; box.line.width = Pt(0.75)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15); tf.margin_top = Inches(0.12); tf.margin_right = Inches(0.15)
    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    r0 = p0.add_run(); r0.text = value
    r0.font.size = Pt(26); r0.font.bold = True; r0.font.name = FONT_MONO; r0.font.color.rgb = INK
    p1 = tf.add_paragraph()
    r1 = p1.add_run(); r1.text = label.upper()
    r1.font.size = Pt(10.5); r1.font.name = FONT_BODY; r1.font.color.rgb = INK_DIM
    return box


def add_bullets(slide, items, left, top, width, height, size=14, bold_lead=True, gap_after=6):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap_after)
        p.line_spacing = 1.15
        r1 = p.add_run(); r1.text = "▸  " + lead
        r1.font.bold = True; r1.font.size = Pt(size); r1.font.name = FONT_BODY; r1.font.color.rgb = INK
        if rest:
            r2 = p.add_run(); r2.text = "  " + rest
            r2.font.size = Pt(size); r2.font.name = FONT_BODY; r2.font.color.rgb = INK_DIM
    return box


def add_status_list(slide, items, left, top, width, height, size=14):
    """items: list of (mark, mark_color, lead, rest)"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (mark, mcolor, lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.line_spacing = 1.15
        rm = p.add_run(); rm.text = mark + "  "
        rm.font.bold = True; rm.font.size = Pt(size); rm.font.name = FONT_MONO; rm.font.color.rgb = mcolor
        r1 = p.add_run(); r1.text = lead
        r1.font.bold = True; r1.font.size = Pt(size); r1.font.name = FONT_BODY; r1.font.color.rgb = INK
        if rest:
            r2 = p.add_run(); r2.text = "  " + rest
            r2.font.size = Pt(size); r2.font.name = FONT_BODY; r2.font.color.rgb = INK_DIM
    return box


def add_tag_row(slide, tags, left, top, size=11):
    x = left
    for t in tags:
        w = Inches(0.24 + 0.11 * len(t))
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top, w, Inches(0.36))
        box.fill.background(); box.line.color.rgb = INK_DIM; box.line.width = Pt(0.75)
        tf = box.text_frame; tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.08)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t.upper()
        r.font.size = Pt(size); r.font.name = FONT_MONO; r.font.bold = True; r.font.color.rgb = INK_DIM
        x = x + w + Inches(0.12)


def add_flow_row(slide, nodes, left, top, node_w, node_h, gap):
    """nodes: list of (title, desc). Draws boxes left-to-right with connecting arrows."""
    x = left
    shapes = []
    for i, (title, desc) in enumerate(nodes):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, top, node_w, node_h)
        box.fill.solid(); box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = ACCENT if len(nodes) > 1 else LINE
        box.line.width = Pt(1)
        tf = box.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.1)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
        r0 = p0.add_run(); r0.text = title.upper()
        r0.font.bold = True; r0.font.size = Pt(12); r0.font.name = FONT_MONO; r0.font.color.rgb = INK
        p1 = tf.add_paragraph(); p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run(); r1.text = desc
        r1.font.size = Pt(9.5); r1.font.name = FONT_BODY; r1.font.color.rgb = INK_DIM
        shapes.append(box)
        x = Emu(int(x) + int(node_w) + int(gap))
        if i < len(nodes) - 1 and gap > Pt(6):
            ax = Emu(int(x) - int(gap))
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, ax, Emu(int(top) + int(node_h)//2 - int(Pt(9))), gap, Pt(18))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = ACCENT; arrow.line.fill.background()
    return shapes


def add_down_arrow(slide, cx, top, size=Pt(20)):
    a = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Emu(int(cx) - int(size)//2), top, size, size)
    a.fill.solid(); a.fill.fore_color.rgb = ACCENT; a.line.fill.background()
    return a


def add_table(slide, left, top, width, height, headers, rows, bad_col=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gshape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = gshape.table
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = SURFACE_2
        tf = cell.text_frame; p = tf.paragraphs[0]
        r = p.add_run(); r.text = h.upper()
        r.font.size = Pt(11); r.font.bold = True; r.font.name = FONT_MONO; r.font.color.rgb = INK_DIM
    for ri, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(ri, c)
            cell.fill.solid(); cell.fill.fore_color.rgb = WHITE
            tf = cell.text_frame; p = tf.paragraphs[0]
            r = p.add_run(); r.text = str(val)
            r.font.size = Pt(12.5)
            r.font.name = FONT_BODY if c == 0 else FONT_MONO
            r.font.color.rgb = BAD if (bad_col is not None and c == bad_col) else INK
            r.font.bold = (bad_col is not None and c == bad_col)
    return table


def add_stamp(slide, text, left, top, color=BAD):
    w, h = Inches(1.5), Inches(0.42)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    box.rotation = -4
    box.fill.background(); box.line.color.rgb = color; box.line.width = Pt(2.25)
    tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text.upper()
    r.font.bold = True; r.font.size = Pt(13); r.font.name = FONT_MONO; r.font.color.rgb = color
    return box


def footer(slide, name):
    add_text(slide, "KAMION CHALLENGE", Inches(0.7), Inches(7.08), Inches(4), Inches(0.3),
              size=9, color=INK_DIM, font=FONT_MONO, upper=True)
    add_text(slide, name, Inches(9.5), Inches(7.08), Inches(3.2), Inches(0.3),
              size=9, color=INK_DIM, font=FONT_MONO, align=PP_ALIGN.RIGHT, upper=True)


# ============================================================ SLIDE 1 — TITLE
s = add_slide(); set_bg(s)
add_text(s, "APPRAISAL SYSTEM — FIELD REPORT", Inches(0.7), Inches(0.5), Inches(6), Inches(0.3),
          size=11, color=INK_DIM, font=FONT_MONO)
add_text(s, "13 SEP 2026", Inches(10.5), Inches(0.5), Inches(2.1), Inches(0.3),
          size=11, color=INK_DIM, font=FONT_MONO, align=PP_ALIGN.RIGHT)
add_eyebrow(s, "Kamion Challenge · 24-Hour Build", top=Inches(2.15))
add_headline(s, "What's This\nTruck Worth?", top=Inches(2.55), size=58, width=Inches(11.5))
add_stripe(s, Inches(4.55))
add_text(s, "A price range, a condition read, and the comparables to back it — computed "
             "from photos alone. No VIN. No mileage. No typed spec sheet.",
          Inches(0.7), Inches(4.75), Inches(9.5), Inches(0.9), size=16, color=INK_DIM)
add_tag_row(s, ["Works", "Sensible", "Knows its limits", "Interesting"], Inches(0.7), Inches(5.9))

# ============================================================ SLIDE 2 — THE BRIEF
s = add_slide(); set_bg(s)
add_eyebrow(s, "The Brief")
add_headline(s, "Photos In. A Checkable\nNumber Out.")
add_stripe(s, Inches(2.15))
add_text(s, "Kamion is a YC-backed freight platform. Today, appraising a used box "
             "truck means a human, a phone camera, and a guess. The challenge: build a "
             "system that looks at the same photos a buyer would and returns something "
             "a judge — or a buyer — can actually check.",
          Inches(0.7), Inches(2.5), Inches(5.6), Inches(3), size=15)
add_status_list(s, [
    ("01", GOOD, "Works", "— live demo, real uploaded photos, no canned output."),
    ("02", GOOD, "Sensible", "— a price range, not a fake point estimate."),
    ("03", GOOD, "Knows its limits", "— rejects bad photos, discloses what it can't see."),
    ("04", GOOD, "Interesting", "— not a thin wrapper around one vision-API call."),
], Inches(6.7), Inches(2.5), Inches(6), Inches(3.2), size=14)
footer(s, "The Brief")

# ============================================================ SLIDE 3 — THE PIPELINE
s = add_slide(); set_bg(s)
add_eyebrow(s, "The Pipeline")
add_headline(s, "Five Stages, Each\nOne Inspectable.")
add_stripe(s, Inches(2.15))
node_w, node_h = Inches(3.6), Inches(0.85)
cx = Inches(6.67)
add_flow_row(s, [("Photos", "1–6 images, any angle")], Emu(int(cx) - int(node_w)//2), Inches(2.5), node_w, node_h, Pt(0))
add_down_arrow(s, cx, Inches(3.42))
add_flow_row(s, [("Gate", "reject dark / blurry / not-a-truck")], Emu(int(cx) - int(node_w)//2), Inches(3.68), node_w, node_h, Pt(0))
add_down_arrow(s, cx, Inches(4.6))
add_flow_row(s, [("CLIP Embedding", "frozen ViT-B/32, mean-pooled")], Emu(int(cx) - int(node_w)//2), Inches(4.86), node_w, node_h, Pt(0))
add_down_arrow(s, cx, Inches(5.78))
fan_w = Inches(3.55)
fan_left = Inches(0.68)
fan_gap = Inches(0.22)
add_flow_row(s, [
    ("Price", "3× quantile LightGBM + CQR"),
    ("Condition", "zero-shot rust / damage / tires"),
    ("Specs", "trained make + weight-class"),
], fan_left, Inches(6.05), fan_w, Inches(0.85), fan_gap)
footer(s, "The Pipeline")

# ============================================================ SLIDE 4 — THE DATA
s = add_slide(); set_bg(s)
add_eyebrow(s, "The Data")
add_headline(s, "877 Real Listings — And\nOne Rejected Shortcut.")
add_stripe(s, Inches(2.15))
add_stat_tile(s, Inches(0.7), Inches(2.45), Inches(1.95), Inches(1.05), "877", "Clean listings")
add_stat_tile(s, Inches(2.75), Inches(2.45), Inches(1.95), Inches(1.05), "~4", "Photos / listing")
add_stat_tile(s, Inches(4.8), Inches(2.45), Inches(2.35), Inches(1.05), "702/86/89", "Train/val/test")
add_text(s, "Scraped from Commercial Truck Trader through a live browser session — its "
             "search data sits behind bot protection that blocks a plain HTTP request. "
             "Deduped, price-outlier-trimmed, and checked photo-by-photo for dealer "
             "placeholder graphics before a single embedding was computed.",
          Inches(0.7), Inches(3.75), Inches(6.4), Inches(2.4), size=14)
add_stamp(s, "Reverted", Inches(10.9), Inches(2.35))
add_text(s, "TRIED: A SECOND MARKETPLACE FOR MORE VOLUME", Inches(7.55), Inches(2.55), Inches(3.1), Inches(0.5),
          size=10.5, color=INK_DIM, font=FONT_MONO)
add_table(s, Inches(7.55), Inches(3.05), Inches(5.1), Inches(1.7),
          ["Metric", "CTT only", "+ 2nd source"],
          [["Price R²", "0.432", "0.294 ↓"],
           ["Make accuracy", "87.6%", "81.8% ↓"],
           ["Class accuracy", "68.5%", "47.6% ↓"]], bad_col=2)
add_text(s, "1 photo/listing on the second source + cross-marketplace noise made every "
             "model worse — confirmed on 5-fold CV, not one lucky split.",
          Inches(7.55), Inches(4.9), Inches(5.1), Inches(1.1), size=11.5, color=INK_DIM)
footer(s, "The Data")

# ============================================================ SLIDE 5 — PRICE MODEL
s = add_slide(); set_bg(s)
add_eyebrow(s, "The Price Model")
add_headline(s, "A Range That's\nActually Calibrated.")
add_stripe(s, Inches(2.15))
add_stat_tile(s, Inches(0.7), Inches(2.5), Inches(2.1), Inches(1.2), "0.432", "Median R² (5-fold CV)")
add_stat_tile(s, Inches(2.95), Inches(2.5), Inches(2.1), Inches(1.2), "$14,191", "Median MAE")
add_stat_tile(s, Inches(5.2), Inches(2.5), Inches(2.1), Inches(1.2), "80.7%", "Coverage (target: 80%)")
add_text(s, "Three LightGBM models predict the 10th / 50th / 90th percentile price from "
             "the image embedding. A held-out calibration split then measures exactly how "
             "wrong the raw 10–90 interval usually is, and widens it by that measured "
             "amount — conformalized quantile regression (CQR).",
          Inches(0.7), Inches(4.05), Inches(6.4), Inches(2), size=14)
box = slide = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.4), Inches(2.5), Inches(5.25), Inches(3.3))
box.fill.solid(); box.fill.fore_color.rgb = SURFACE_2; box.line.color.rgb = LINE
tf = box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.25)
p0 = tf.paragraphs[0]; r0 = p0.add_run(); r0.text = "WHY THIS MATTERS"
r0.font.bold = True; r0.font.size = Pt(11); r0.font.name = FONT_MONO; r0.font.color.rgb = INK_DIM
p1 = tf.add_paragraph(); p1.space_before = Pt(10)
r1 = p1.add_run(); r1.text = ("“80% confidence” is a claim you can audit: it means on "
                                "held-out listings the real price landed inside the printed "
                                "range 80.7% of the time — measured, not asserted.")
r1.font.size = Pt(15); r1.font.name = FONT_BODY; r1.font.color.rgb = INK
footer(s, "The Price Model")

# ============================================================ SLIDE 6 — THE HONEST CEILING
s = add_slide(); set_bg(s)
add_eyebrow(s, "The Honest Ceiling")
add_headline(s, "The Camera Can't\nSee The Odometer.")
add_stripe(s, Inches(2.15))
chart_data = CategoryChartData()
chart_data.categories = ["Full photo set", "Mileage + year alone"]
chart_data.add_series("R² (5-fold CV)", (0.432, 0.602))
gframe = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.7), Inches(2.5), Inches(6.1), Inches(4.3), chart_data)
chart = gframe.chart
chart.has_legend = False
plot = chart.plots[0]
plot.has_data_labels = True
plot.data_labels.number_format = "0.00"
plot.data_labels.number_format_is_linked = False
plot.data_labels.font.size = Pt(14)
plot.data_labels.font.bold = True
plot.data_labels.font.color.rgb = INK
series = plot.series[0]
series.points[0].format.fill.solid(); series.points[0].format.fill.fore_color.rgb = INK_DIM
series.points[1].format.fill.solid(); series.points[1].format.fill.fore_color.rgb = ACCENT
cat_ax = chart.category_axis
cat_ax.tick_labels.font.size = Pt(12); cat_ax.tick_labels.font.color.rgb = INK
val_ax = chart.value_axis
val_ax.minimum_scale = 0; val_ax.maximum_scale = 1.0
val_ax.tick_labels.font.size = Pt(10); val_ax.tick_labels.font.color.rgb = INK_DIM
add_text(s, "Two numbers no camera can see — mileage and model year — explain "
             "more of a truck's price (R²=0.60) than the full image embedding does "
             "(R²=0.43).",
          Inches(7.1), Inches(2.6), Inches(5.6), Inches(1.6), size=15)
add_text(s, "We tested three fixes for a too-wide interval — loosening the tail "
             "models, a different calibration method, feeding in predicted vehicle class "
             "— none closed this gap. It's structural, not a bug: this is the measured "
             "floor under how narrow a photo-only range can honestly get.",
          Inches(7.1), Inches(4.3), Inches(5.6), Inches(2.3), size=13.5, color=INK_DIM)
footer(s, "The Honest Ceiling")

# ============================================================ SLIDE 7 — CONDITION, VERIFIED
s = add_slide(); set_bg(s)
add_eyebrow(s, "Condition — Verified, Not Assumed")
add_headline(s, "One Flag We Trust.\nThree We Don't — Yet.")
add_stripe(s, Inches(2.15))
add_text(s, "Zero-shot CLIP flags rust, body damage, tire wear, and interior wear "
             "against the fleet's own score distribution. The first threshold (80th "
             "percentile) looked reasonable — until we hand-labeled 39 real trucks "
             "and found it flagging 10 of them for “worn tires” with nothing wrong.\n\n"
             "Recalibrated to the 98th percentile: 0 false positives on the same 39 "
             "trucks. Then scaled the ground truth to 113 hand-labeled listings across 3 "
             "targeted rounds to check every attribute, not just the easy ones.",
          Inches(0.7), Inches(2.5), Inches(6.3), Inches(4), size=14)
add_status_list(s, [
    ("✓", GOOD, "Rust", "— caught both real corrosion cases found in testing. Shown live."),
    ("✗", BAD, "Body damage", "— missed a real defect (graffiti). Experimental."),
    ("✗", BAD, "Tire wear", "— missed its one real case found. Experimental."),
    ("✗", BAD, "Interior wear", "— missed both real cases found. Experimental."),
], Inches(7.35), Inches(2.6), Inches(5.3), Inches(4), size=14)
footer(s, "Condition, Verified")

# ============================================================ SLIDE 8 — SPECS FROM PIXELS
s = add_slide(); set_bg(s)
add_eyebrow(s, "Specs From Pixels")
add_headline(s, "A Ford That Looked\nLike A Chevy.")
add_stripe(s, Inches(2.15))
add_stat_tile(s, Inches(0.7), Inches(2.5), Inches(2.9), Inches(1.2), "87.6%", "Make accuracy (trained classifier)")
add_stat_tile(s, Inches(3.75), Inches(2.5), Inches(2.9), Inches(1.2), "68.5%", "Weight-class accuracy")
add_text(s, "Both are logistic regression trained on the same frozen embeddings — "
             "not a nearest-neighbor guess.",
          Inches(0.7), Inches(3.95), Inches(6), Inches(0.9), size=13.5, color=INK_DIM)
box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.05), Inches(2.5), Inches(5.6), Inches(4.2))
box.fill.solid(); box.fill.fore_color.rgb = SURFACE_2; box.line.color.rgb = LINE
tf = box.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.22); tf.margin_right = Inches(0.25)
p0 = tf.paragraphs[0]; r0 = p0.add_run(); r0.text = "LIVE FAILURE, ROOT-CAUSED"
r0.font.bold = True; r0.font.size = Pt(11); r0.font.name = FONT_MONO; r0.font.color.rgb = INK_DIM
p1 = tf.add_paragraph(); p1.space_before = Pt(10)
r1 = p1.add_run(); r1.text = ("A real 2026 Ford F-750 came back “Chevrolet, 49% "
                                "confidence.” Cause: “Ford” in the training data is 53% "
                                "Econoline cutaway vans and only ~17% heavy conventional "
                                "trucks — two unrelated body shapes sharing one badge.")
r1.font.size = Pt(14); r1.font.name = FONT_BODY; r1.font.color.rgb = INK
p2 = tf.add_paragraph(); p2.space_before = Pt(12)
r2 = p2.add_run(); r2.text = ("Fix: added the GVWR weight-class model as a "
                                "visually-consistent second signal, shown alongside make "
                                "— softer errors, size instead of badge.")
r2.font.size = Pt(14); r2.font.name = FONT_BODY; r2.font.color.rgb = INK
footer(s, "Specs From Pixels")

# ============================================================ SLIDE 9 — THE LIVE DEMO
s = add_slide(); set_bg(s)
add_eyebrow(s, "The Live Demo")
add_headline(s, "Upload A Photo,\nGet An Appraisal.")
add_stripe(s, Inches(2.15))
add_flow_row(s, [
    ("1 · Upload", "any angle, any count"),
    ("2 · Price", "range + confidence badge"),
    ("3 · Specs", "weight class + make"),
], Inches(0.68), Inches(2.6), Inches(3.85), Inches(1.0), Inches(0.18))
add_flow_row(s, [
    ("4 · Condition", "verified vs. experimental flags"),
    ("5 · Why This Price", "3 nearest real comparables"),
], Inches(0.68), Inches(3.9), Inches(5.9), Inches(1.0), Inches(0.18))
add_text(s, "A dark, blurry, or non-truck photo doesn't get a confident wrong answer "
             "— it gets rejected with a specific, actionable reason.",
          Inches(0.7), Inches(5.4), Inches(9.5), Inches(0.9), size=15, color=INK_DIM)
footer(s, "The Live Demo")

# ============================================================ SLIDE 10 — WHAT'S NEXT
s = add_slide(); set_bg(s)
add_eyebrow(s, "What's Next")
add_headline(s, "Still To Do.")
add_stripe(s, Inches(1.85))
add_bullets(s, [
    ("Broaden the body-damage prompt", "to catch cosmetic defects like graffiti, not just dents."),
    ("Source higher-variance listings", "(private-party, as-is) to properly test recall on real damage."),
    ("More scraped volume", "— the honest way to legitimately narrow the price range further."),
], Inches(0.7), Inches(2.35), Inches(6.2), Inches(3), size=15, gap_after=14)
add_text(s, "What's This\nTruck Worth?", Inches(7.3), Inches(2.6), Inches(5.3), Inches(1.6),
          size=34, color=INK, bold=True, font=FONT_HEAD, align=PP_ALIGN.CENTER)
add_text(s, "Built in 24 hours for the Kamion Challenge.", Inches(7.3), Inches(4.15), Inches(5.3), Inches(0.6),
          size=13, color=INK_DIM, align=PP_ALIGN.CENTER)
footer(s, "What's Next")

out_path = "Kamion_Challenge_Presentation.pptx"
prs.save(out_path)
print(f"Wrote {out_path}")
